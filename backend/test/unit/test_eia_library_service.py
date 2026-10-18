"""eia_library_service 单测：sha1 命名/去重、越限拒收、移动、删除、清单过滤（仓储以内存桩替身）。

P2 增量：MinIO 正源（FakeMinIO 替身）、canonical md 同步产出（ready/failed/not_applicable）、
report_id 自动生成、update_meta 值域守卫、get_canonical_md 往返。
"""

from __future__ import annotations

import asyncio
import io
from types import SimpleNamespace

import pytest
from fastapi import UploadFile

from yuxi.services.eia_library_service import (
    EiaLibraryService,
    LibraryFileTooLarge,
    _safe_filename,
)

SHA1_HELLO = "aaf4c61ddcc5e8a2dabede0f3b482cd9aea9434d"  # sha1(b"hello")


class FakeRepo:
    def __init__(self):
        self.folders = {
            1: SimpleNamespace(id=1, name="规划环评", stage_key="planning_eia"),
            2: SimpleNamespace(id=2, name="后评价", stage_key="post_eia"),
        }
        self.files: list = []
        self._next = 1

    async def get_folder(self, folder_id):
        return self.folders.get(folder_id)

    async def get_file_by_sha1(self, sha1):
        return next((f for f in self.files if f.content_sha1 == sha1), None)

    async def list_files(self, *, folder_id=None, q=None, uncategorized=False):
        rows = self.files
        if folder_id is not None:
            rows = [f for f in rows if f.folder_id == folder_id]
        elif uncategorized:
            rows = [f for f in rows if f.folder_id is None]
        if q:
            rows = [f for f in rows if q.lower() in f.original_filename.lower()]
        return [
            {**f.to_dict(), "folder_name": self.folders[f.folder_id].name if f.folder_id in self.folders else None}
            for f in rows
        ]

    async def insert_file(self, row):
        row.id = self._next
        self._next += 1
        self.files.append(row)
        return row

    async def update_file(self, file_id, fields):
        row = await self.get_file(file_id)
        if row is None:
            return None
        for k, v in fields.items():
            setattr(row, k, v)
        return row

    async def get_file(self, file_id):
        return next((f for f in self.files if f.id == file_id), None)

    async def move_file(self, file_id, folder_id):
        row = await self.get_file(file_id)
        if row is not None:
            row.folder_id = folder_id
        return row

    async def delete_file(self, file_id):
        row = await self.get_file(file_id)
        if row is None:
            return False
        self.files.remove(row)
        return True


class FakeUploadResult:
    def __init__(self, url):
        self.url = url


class FakeMinIO:
    """内存桶替身：接口对齐 yuxi.storage.minio.MinIOClient 异步面。"""

    def __init__(self):
        self.store: dict[str, bytes] = {}

    async def aupload_file(self, bucket, obj, data, content_type=None):
        self.store[obj] = bytes(data)
        return FakeUploadResult(f"http://minio:9000/{bucket}/{obj}")

    async def adownload_file(self, bucket, obj):
        if obj not in self.store:
            raise FileNotFoundError(obj)
        return self.store[obj]

    async def adelete_file(self, bucket, obj):
        return self.store.pop(obj, None) is not None

    async def astat_file(self, bucket, obj):
        return len(self.store[obj]) if obj in self.store else None


def _docx_bytes(cover: list[str], body: list[str]) -> bytes:
    """最小合成 docx：封面段 + 正文节号标题段。"""
    from docx import Document

    doc = Document()
    for t in cover:
        doc.add_paragraph(t)
    for t in body:
        doc.add_paragraph(t)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


GOOD_DOCX = _docx_bytes(["某矿项目环境影响报告书", "工程编号：H1"], ["1 概述", "内容。"])


@pytest.fixture(autouse=True)
def fake_minio(monkeypatch):
    store = FakeMinIO()
    monkeypatch.setattr("yuxi.services.eia_library_service.get_minio_client", lambda: store)
    return store


def make_service(tmp_path, max_size=None):
    return EiaLibraryService(repo=FakeRepo(), storage_root=tmp_path, max_size_bytes=max_size)


def make_upload(data: bytes, filename: str) -> UploadFile:
    return UploadFile(file=io.BytesIO(data), filename=filename)


def test_upload_stores_sha1_named_object_in_minio(tmp_path, fake_minio):
    svc = make_service(tmp_path)
    row = asyncio.run(svc.upload(make_upload(b"hello", "某矿区报告书.docx"), folder_id=1, uploaded_by="admin"))
    assert row["content_sha1"] == SHA1_HELLO
    assert row["stored_path"] == f"{SHA1_HELLO[:12]}__某矿区报告书.docx"
    assert row["object_key"] == f"eia-library/{SHA1_HELLO[:12]}__某矿区报告书.docx"
    assert row["size_bytes"] == 5
    assert fake_minio.store[row["object_key"]] == b"hello"  # 正源在 MinIO
    assert not (tmp_path / row["stored_path"]).exists()  # 本地不再落副本
    assert row["status"] == "stored"


def test_upload_docx_produces_canonical_md_and_rid(tmp_path, fake_minio):
    svc = make_service(tmp_path)
    row = asyncio.run(svc.upload(make_upload(GOOD_DOCX, "样例报告.docx"), folder_id=1, uploaded_by="tester"))
    assert row["md_status"] == "ready"
    assert row["md_chars"] > 0
    assert row["md_sha1"]
    assert row["report_id"].startswith("001-")  # make_rid(id, stem)
    assert row["stage"] == "planning_eia"  # 文件夹 stage_key 作默认
    md_obj = f"eia-library/.canonical/{row['stored_path']}.md"
    assert fake_minio.store[md_obj].decode("utf-8").startswith("某矿项目环境影响报告书")  # 前置区
    assert "<!-- EIA-FRONT-MATTER-END -->" in fake_minio.store[md_obj].decode("utf-8")
    assert row["front_meta"]["project_no"] == "H1"


def test_upload_non_docx_not_applicable(tmp_path, fake_minio):
    svc = make_service(tmp_path)
    row = asyncio.run(svc.upload(make_upload(b"%PDF-1.4", "说明.pdf"), folder_id=None, uploaded_by="tester"))
    assert row["md_status"] == "not_applicable"
    assert row["md_object_key"] is None


def test_upload_md_failure_keeps_registration(tmp_path, fake_minio):
    svc = make_service(tmp_path)
    row = asyncio.run(svc.upload(make_upload(b"hello", "坏文档.docx"), folder_id=None, uploaded_by="tester"))
    assert row["md_status"] == "failed"
    assert row["md_failed_reason"]  # 伪 docx（corrupt）→ failed + 原因，登记照记
    assert row["report_id"].startswith("001-")


def test_upload_duplicate_sha1_rejected_and_tmp_cleaned(tmp_path, fake_minio):
    svc = make_service(tmp_path)
    asyncio.run(svc.upload(make_upload(b"hello", "a.docx"), folder_id=None, uploaded_by="admin"))
    with pytest.raises(FileExistsError):
        asyncio.run(svc.upload(make_upload(b"hello", "b.docx"), folder_id=None, uploaded_by="admin"))
    assert len(svc.repo.files) == 1
    assert not [p for p in tmp_path.iterdir() if p.name.startswith(".tmp-")]


def test_upload_oversize_raises_too_large(tmp_path, fake_minio):
    svc = make_service(tmp_path, max_size=4)
    with pytest.raises(LibraryFileTooLarge):
        asyncio.run(svc.upload(make_upload(b"0123456789", "big.docx"), folder_id=None, uploaded_by="admin"))
    assert svc.repo.files == []


def test_upload_to_unknown_folder_rejected(tmp_path, fake_minio):
    svc = make_service(tmp_path)
    with pytest.raises(ValueError, match="目标文件夹不存在"):
        asyncio.run(svc.upload(make_upload(b"x", "a.docx"), folder_id=99, uploaded_by="admin"))


def test_move_updates_folder(tmp_path, fake_minio):
    svc = make_service(tmp_path)
    row = asyncio.run(svc.upload(make_upload(b"hello", "a.docx"), folder_id=1, uploaded_by="admin"))
    out = asyncio.run(svc.move(row["id"], 2))
    assert out["folder_id"] == 2


def test_delete_removes_row_and_minio_objects(tmp_path, fake_minio):
    svc = make_service(tmp_path)
    row = asyncio.run(svc.upload(make_upload(GOOD_DOCX, "a.docx"), folder_id=None, uploaded_by="admin"))
    src_obj = row["object_key"]
    md_obj = row["md_object_key"]
    assert src_obj in fake_minio.store and md_obj in fake_minio.store
    asyncio.run(svc.delete(row["id"]))
    assert svc.repo.files == []
    assert src_obj not in fake_minio.store and md_obj not in fake_minio.store
    with pytest.raises(FileNotFoundError):
        asyncio.run(svc.delete(row["id"]))


def test_get_canonical_md_roundtrip(tmp_path, fake_minio):
    svc = make_service(tmp_path)
    row = asyncio.run(svc.upload(make_upload(GOOD_DOCX, "a.docx"), folder_id=None, uploaded_by="admin"))
    got_row, data = asyncio.run(svc.get_canonical_md(row["id"]))
    assert got_row.id == row["id"]
    assert "<!-- EIA-FRONT-MATTER-END -->" in data.decode("utf-8")
    with pytest.raises(FileNotFoundError):
        asyncio.run(svc.get_canonical_md(999))


def test_regenerate_md_recovers_failed_row(tmp_path, fake_minio):
    svc = make_service(tmp_path)
    broken = asyncio.run(svc.upload(make_upload(b"hello", "b.docx"), folder_id=None, uploaded_by="admin"))
    assert broken["md_status"] == "failed"
    # 同名真 docx 字节重产：直接以 MinIO 正源替换后调用
    fake_minio.store[broken["object_key"]] = GOOD_DOCX
    out = asyncio.run(svc.regenerate_md(broken["id"], "tester"))
    assert out["md_status"] == "ready"
    assert out["md_chars"] > 0


def test_update_meta_stage_whitelist(tmp_path, fake_minio):
    svc = make_service(tmp_path)
    row = asyncio.run(svc.upload(make_upload(b"hello", "a.docx"), folder_id=None, uploaded_by="admin"))
    out = asyncio.run(svc.update_meta(row["id"], {"stage": "post_eia", "mineral": "coal"}))
    assert out["stage"] == "post_eia"
    with pytest.raises(ValueError, match="非法 stage"):
        asyncio.run(svc.update_meta(row["id"], {"stage": "not_a_stage"}))


def test_list_files_joins_folder_name(tmp_path, fake_minio):
    svc = make_service(tmp_path)
    asyncio.run(svc.upload(make_upload(b"hello", "alpha.docx"), folder_id=1, uploaded_by="admin"))
    asyncio.run(svc.upload(make_upload(b"world", "beta.docx"), folder_id=None, uploaded_by="admin"))
    all_files = asyncio.run(svc.list_files())
    assert len(all_files) == 2
    names = {f["original_filename"]: f["folder_name"] for f in all_files}
    assert names["alpha.docx"] == "规划环评"
    assert names["beta.docx"] is None
    hit = asyncio.run(svc.list_files(q="alp"))
    assert [f["original_filename"] for f in hit] == ["alpha.docx"]


def test_list_files_uncategorized_filters_null_folder(tmp_path, fake_minio):
    svc = make_service(tmp_path)
    asyncio.run(svc.upload(make_upload(b"hello", "alpha.docx"), folder_id=1, uploaded_by="admin"))
    asyncio.run(svc.upload(make_upload(b"world", "beta.docx"), folder_id=None, uploaded_by="admin"))
    hit = asyncio.run(svc.list_files(uncategorized=True))
    assert [f["original_filename"] for f in hit] == ["beta.docx"]


def test_safe_filename_strips_unsafe_chars():
    assert _safe_filename(r"report\v2:final?.docx") == "report_v2_final_.docx"
    assert _safe_filename("..") == "unnamed"
